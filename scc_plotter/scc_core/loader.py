"""Read the SCC crack-growth-rate workbook into a clean pandas DataFrame.

The workbook has a few quirks this module handles so the plotting code does not
have to:

* The real header row is not row 1. Above it sit a title, a highlighting legend
  and a row of explanatory notes ("compact tension id", "tip of the crack", ...).
* Each header cell holds "name<newline>unit", e.g. "CGR\\nmm/s".
* Some rows are highlighted (green fill in column A = "PWR Mainstream Data").
  The meaning of each highlight is written in the legend at the top of the sheet,
  e.g. "PWR Mainstream Data in Column A (265 data)". Those highlights become
  True/False flag columns.
* Numeric columns contain stray text such as "－" for "not reported".
"""

from __future__ import annotations

import re
from dataclasses import dataclass, field

import openpyxl
import pandas as pd
from openpyxl.utils import column_index_from_string

# A column is treated as numeric when at least this share of its non-empty
# cells parse as numbers. The rest (e.g. "－") become empty.
NUMERIC_SHARE = 0.5

# Columns that look numeric but are really labels (material grade 304 / 316 ...).
FORCE_CATEGORICAL = {"Material", "Heat", "CT / ID", "Lab", "Condition"}

LEGEND_RE = re.compile(r"^(?P<label>.+?)\s+in\s+Column\s+(?P<col>[A-Z]{1,3})\b", re.I)


@dataclass
class ColumnInfo:
    header: str          # original header text, single line
    name: str            # display name, e.g. "CGR"
    unit: str            # unit text, e.g. "mm/s" ("" when none)
    var: str             # identifier usable in formulas, e.g. "CGR"
    kind: str            # "numeric", "categorical" or "flag"
    note: str = ""       # explanatory note found above the header
    source: str = "excel"  # "excel", "highlight" or "custom"
    n_unparsed: int = 0  # text cells turned into empty in a numeric column

    @property
    def label(self) -> str:
        return f"{self.name} [{self.unit}]" if self.unit else self.name


@dataclass
class Dataset:
    df: pd.DataFrame
    columns: dict[str, ColumnInfo]  # keyed by DataFrame column name (== var)
    sheet: str
    header_row: int                 # 1-based Excel row number
    warnings: list[str] = field(default_factory=list)

    def numeric_vars(self) -> list[str]:
        return [v for v, c in self.columns.items() if c.kind == "numeric"]

    def categorical_vars(self) -> list[str]:
        return [v for v, c in self.columns.items() if c.kind in ("categorical", "flag")]


def make_var(name: str) -> str:
    """Turn a header name into a Python identifier: '∆Time' -> 'dTime'."""
    s = name.replace("∆", "d").replace("Δ", "d").replace("%", "")
    s = re.sub(r"[^0-9A-Za-z_]+", "_", s).strip("_")
    if not s:
        s = "col"
    if s[0].isdigit():
        s = "c_" + s
    return s


def split_header(text: str) -> tuple[str, str]:
    """'CGR\\nmm/s' -> ('CGR', 'mm/s'); '%CW' -> ('%CW', '%')."""
    parts = [p.strip() for p in str(text).splitlines() if p.strip()]
    name = parts[0] if parts else str(text).strip()
    unit = " ".join(parts[1:])
    if not unit and name.startswith("%"):
        unit = "%"
    return name, unit


def _find_header_row(rows: list[list], max_scan: int = 20) -> int:
    """Index of the row with the most text cells among the first rows.

    Ties go to the lower row, since the notes row sits above the header.
    """
    best, best_score = 0, -1
    for i, row in enumerate(rows[:max_scan]):
        score = sum(isinstance(v, str) and v.strip() != "" for v in row)
        if score >= best_score:
            best, best_score = i, score
    return best


def _pick_sheet(wb: openpyxl.Workbook, sheet: str | None) -> str:
    if sheet:
        return sheet
    # The data sheet is the biggest one ("Sheet1" holds only a short excerpt).
    return max(wb.sheetnames, key=lambda n: wb[n].max_row * wb[n].max_column)


def load_workbook(path, sheet: str | None = None) -> Dataset:
    """Load the workbook at `path` (a path or a file-like object)."""
    wb = openpyxl.load_workbook(path, data_only=True)
    sheet = _pick_sheet(wb, sheet)
    ws = wb[sheet]

    rows = [list(r) for r in ws.iter_rows(values_only=True)]
    h = _find_header_row(rows)
    header = rows[h]
    notes = rows[h - 1] if h > 0 else [None] * len(header)

    # Keep only columns that have a header.
    col_idx = [j for j, v in enumerate(header) if v is not None and str(v).strip()]
    data = [[r[j] if j < len(r) else None for j in col_idx] for r in rows[h + 1:]]
    raw = pd.DataFrame(data, dtype=object)
    # Drop fully empty rows (trailing blank lines in the sheet).
    keep = raw.apply(lambda r: any(v is not None and str(v).strip() != "" for v in r), axis=1)
    raw = raw[keep]
    excel_rows = [h + 2 + i for i in raw.index]  # 1-based Excel row numbers
    raw = raw.reset_index(drop=True)

    columns: dict[str, ColumnInfo] = {}
    out = {}
    warnings: list[str] = []
    for k, j in enumerate(col_idx):
        name, unit = split_header(header[j])
        var = make_var(name)
        while var in columns:
            var += "_"
        note = notes[j] if j < len(notes) and isinstance(notes[j], str) else ""
        series = raw[k].map(_blank_to_none)

        numeric = pd.to_numeric(series, errors="coerce")
        n_present = int(series.notna().sum())
        n_numeric = int(numeric.notna().sum())
        is_numeric = (
            name not in FORCE_CATEGORICAL
            and n_present > 0
            and n_numeric / n_present >= NUMERIC_SHARE
        )
        if is_numeric:
            out[var] = numeric.astype(float)
            n_bad = n_present - n_numeric
            if n_bad:
                bad = sorted({str(v) for v in series[numeric.isna() & series.notna()]})[:5]
                warnings.append(
                    f"{name}: {n_bad} non-numeric cell(s) treated as empty (e.g. {', '.join(bad)})"
                )
            kind = "numeric"
        else:
            out[var] = series.map(lambda v: None if _blank_to_none(v) is None else str(v).strip()).astype("string")
            n_bad = 0
            kind = "categorical"
        columns[var] = ColumnInfo(
            header=" ".join(str(header[j]).split()), name=name, unit=unit, var=var,
            kind=kind, note=note.strip(), n_unparsed=n_bad,
        )

    df = pd.DataFrame(out)

    # Highlight flags described by the legend above the header.
    for r in rows[:h]:
        for v in r:
            m = LEGEND_RE.match(str(v).strip()) if isinstance(v, str) else None
            if not m:
                continue
            letter = m.group("col").upper()
            cidx = column_index_from_string(letter)
            flags = [_has_fill(ws.cell(row=er, column=cidx)) for er in excel_rows]
            label = m.group("label").strip()
            var = make_var(label)
            df[var] = flags
            columns[var] = ColumnInfo(
                header=label, name=label, unit="", var=var, kind="flag",
                note=f"Rows highlighted in column {letter} of the workbook", source="highlight",
            )

    df.insert(0, "excel_row", excel_rows)
    return Dataset(df=df, columns=columns, sheet=sheet, header_row=h + 1, warnings=warnings)


def _blank_to_none(v):
    if v is None or (isinstance(v, float) and pd.isna(v)):
        return None
    if isinstance(v, str) and (not v.strip() or v.strip() == "_x000D_"):
        return None
    return v


def _has_fill(cell) -> bool:
    fill = cell.fill
    if fill is None or fill.fill_type in (None, "none"):
        return False
    color = fill.fgColor
    # Theme colours have no rgb string; any solid theme fill counts as highlight.
    if color is not None and color.type == "rgb":
        rgb = color.rgb if isinstance(color.rgb, str) else ""
        return rgb not in ("00000000", "FFFFFFFF", "")
    return True
