"""Tests for the data and formula logic. Run: python -m pytest tests  (or python tests/test_core.py)

Uses a small synthetic workbook that mimics the real one's layout, so the
tests do not need the (non-public) database file.
"""

import io
import sys
from pathlib import Path

import numpy as np
import openpyxl
import pandas as pd
from openpyxl.styles import PatternFill

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

from scc_core.custom_params import (  # noqa: E402
    CustomParam, FormulaError, apply_params, check_formula, load_params, save_params,
)
from scc_core.loader import load_workbook, make_var, split_header  # noqa: E402
from scc_core.plotdata import (  # noqa: E402
    EMPTY, OTHER, Filters, apply_filters, category_colors, coverage, prepare_xy,
)


def make_book() -> io.BytesIO:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "DB"
    ws["A1"] = "Overall database"
    ws["G1"] = "With highlighting for:"
    ws["I1"] = "Mainstream Data in Column A (2 data)"
    ws["I1"].fill = PatternFill("solid", fgColor="FF92D050")
    ws["C3"] = "compact tension id"
    header = ["Index", "Lab", "CT / ID", "Material", "%CW", "∆Time\nh", "∆a\nmm",
              "CGR\nmm/s", "K\nMPa√m", "Temp.\n°C"]
    for j, h in enumerate(header, start=1):
        ws.cell(row=4, column=j, value=h)
    rows = [
        [1, "GE", "C1", 304, 20, 1000, 0.36, 1e-7, 30, 320],
        [2, "GE", "C1", "316L", 10, "－", "－", 2e-8, 25, 290],
        [3, "INSS", None, "316L", 0, 500, 0.0, 0.0, 40, 290],
        [4, "MHI", "M7", "316L", 30, 200, 0.072, None, 20, None],
    ]
    for i, r in enumerate(rows, start=5):
        for j, v in enumerate(r, start=1):
            ws.cell(row=i, column=j, value=v)
    for r in (5, 7):
        ws.cell(row=r, column=1).fill = PatternFill("solid", fgColor="FF92D050")
    wb.create_sheet("Sheet1")["A1"] = "excerpt"
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf


def test_header_helpers():
    assert split_header("CGR\nmm/s") == ("CGR", "mm/s")
    assert split_header("%CW") == ("%CW", "%")
    assert split_header("H2 \ncc/kg") == ("H2", "cc/kg")
    assert make_var("∆Time") == "dTime"
    assert make_var("CT / ID") == "CT_ID"
    assert make_var("Temp.") == "Temp"
    assert make_var("%CW") == "CW"


def test_load_workbook():
    ds = load_workbook(make_book())
    assert ds.sheet == "DB"            # biggest sheet, not the excerpt
    assert ds.header_row == 4
    df = ds.df
    assert len(df) == 4
    assert ds.columns["CGR"].unit == "mm/s"
    assert ds.columns["CT_ID"].note == "compact tension id"
    assert ds.columns["Material"].kind == "categorical"
    assert list(df["Material"]) == ["304", "316L", "316L", "316L"]
    assert np.isnan(df.loc[1, "dTime"])  # "－" became empty
    assert any("∆a" in w for w in ds.warnings)
    assert list(df["Mainstream_Data"]) == [True, False, True, False]
    assert ds.columns["Mainstream_Data"].kind == "flag"
    assert list(df["excel_row"]) == [5, 6, 7, 8]


def test_formulas_and_lookups():
    ds = load_workbook(make_book())
    params = [
        CustomParam(name="R", kind="constant", value=8.314),
        CustomParam(name="T_K", kind="formula", formula="Temp + 273.15"),
        CustomParam(name="invT", kind="formula", formula="1000 / T_K"),
        CustomParam(name="rate", kind="formula", formula="da / (dTime * 3600)"),
        CustomParam(name="E", kind="lookup", key_column="Material",
                    mapping={"304": 193.0}, default=200.0),
        CustomParam(name="bad", kind="formula", formula="__import__('os')"),
    ]
    out, errors = apply_params(ds, params)
    df = out.df
    assert np.isclose(df.loc[0, "invT"], 1000 / 593.15)
    assert np.isnan(df.loc[3, "invT"])
    assert np.isclose(df.loc[0, "rate"], 1e-7)
    assert list(df["E"]) == [193.0, 200.0, 200.0, 200.0]
    assert len(errors) == 1 and errors[0].startswith("bad")
    assert out.columns["T_K"].source == "custom"


def test_check_formula_rejects_unsafe():
    for bad in ["x.__class__", "open('f')", "[1, 2]", "lambda: 1", "K if K else 1"]:
        try:
            check_formula(bad, {"K", "x"})
        except FormulaError:
            continue
        raise AssertionError(f"accepted {bad!r}")
    check_formula("log10(K) * 2 + sqrt(abs(K))", {"K"})


def test_params_roundtrip(tmp_path=None):
    path = Path(tmp_path or ".") / "_params_test.json"
    try:
        p = [CustomParam(name="E", kind="lookup", key_column="Material", mapping={"304": 1.0})]
        save_params(path, p)
        assert load_params(path) == p
    finally:
        path.unlink(missing_ok=True)


def test_shipped_params_file_loads():
    path = Path(__file__).resolve().parents[1] / "custom_parameters.json"
    ds = load_workbook(make_book())
    out, errors = apply_params(ds, load_params(path))
    assert errors == []
    assert {"T_K", "invT_1000", "CGR_from_da"} <= set(out.df.columns)


def test_prepare_xy_and_filters():
    ds = load_workbook(make_book())
    df = ds.df
    res = prepare_xy(df, "K", "CGR", False, True)
    assert len(res.data) == 2 and res.n_missing_y == 1 and res.n_nonpositive == 1
    f = Filters(categories={"Lab": ["GE"]})
    assert len(apply_filters(df, f)) == 2
    f = Filters(ranges={"Temp": (300, 400)}, flags=["Mainstream_Data"])
    assert list(apply_filters(df, f)["Index"]) == [1]
    f = Filters(categories={"CT_ID": [EMPTY]})
    assert list(apply_filters(df, f)["Index"]) == [3]


def test_category_colors_stable_and_folded():
    s = pd.Series(["a"] * 5 + ["b"] * 3 + ["c"] * 2 + ["d"] + [None])
    display, cmap, order = category_colors(s, max_colors=2)
    assert order == ["a", "b", OTHER, EMPTY]
    assert display.iloc[-1] == EMPTY and display.iloc[-2] == OTHER
    # Colours come from the full data, so a filtered view keeps them.
    assert cmap["a"] != cmap["b"]


def test_coverage():
    ds = load_workbook(make_book())
    cov = coverage(ds.df, "Lab", ["dTime", "CGR"])
    assert cov.loc["GE", "dTime"] == 50.0
    assert cov.loc["GE", "rows"] == 2


def test_units():
    from scc_core.units import parse_unit, UnitError
    u = parse_unit("MPa√m")
    assert np.isclose(u.scale, 1e6) and np.allclose(u.dims, (1, -0.5, -2, 0, 0))
    assert np.isclose(parse_unit("mm/s").scale, 1e-3)
    assert parse_unit("°C").offset == 273.15
    assert np.allclose(parse_unit("J/(mol·K)").dims, (1, 2, -2, -1, -1))
    assert np.isclose(parse_unit("m²/s").scale, 1) and np.allclose(parse_unit("m²/s").dims, (0, 2, -1, 0, 0))
    assert np.isclose(parse_unit("µm").scale, 1e-6)
    assert parse_unit("ppb").dimensionless and parse_unit("").dimensionless
    try:
        parse_unit("1 to 5")
        raise AssertionError("accepted '1 to 5'")
    except UnitError:
        pass


def test_buckingham_match1_groups():
    from fractions import Fraction
    from scc_core.units import parse_unit
    from scc_core.dimensional import buckingham
    from scc_core.matches import MATCH_1
    units = {"CGR": parse_unit("mm/s"), "K": parse_unit("MPa√m"), "T_K": parse_unit("K")}
    units.update({p.name: parse_unit(p.unit) for p in MATCH_1["params"]})
    r = buckingham(MATCH_1["variables"], units, MATCH_1["repeating"])
    assert r.rank == 5 and r.n_groups == 3
    g = {x.main: x.exponents for x in r.groups}
    assert g["CGR"] == {"CGR": 1, "d": 1, "D": -1}
    assert g["K"] == {"K": 1, "sigma_y": -1, "d": Fraction(-1, 2)}
    assert g["Q"] == {"Q": 1, "R": -1, "T_K": -1}


def test_group_value_in_si():
    from scc_core.units import parse_unit
    from scc_core.dimensional import buckingham
    from scc_core.custom_params import evaluate_formula
    df = pd.DataFrame({"CGR": [1e-6], "K": [30.0], "Temp": [290.0]})
    units = {"CGR": parse_unit("mm/s"), "K": parse_unit("MPa√m"), "Temp": parse_unit("°C"),
             "s_y": parse_unit("MPa"), "d": parse_unit("µm")}
    r = buckingham(["K", "s_y", "d"], units, ["s_y", "d"])
    val = evaluate_formula(r.groups[0].formula(units), df, {"s_y": 400.0, "d": 50.0})[0]
    assert np.isclose(val, 30e6 / (400e6 * np.sqrt(50e-6)))


def test_formula_dims_and_si_rewrite():
    from scc_core.units import parse_unit
    from scc_core.dimensional import formula_dims, is_dimensionless, si_formula, DimensionError
    units = {"Temp": parse_unit("°C"), "Q": parse_unit("J/mol"), "R": parse_unit("J/(mol·K)"),
             "T_K": parse_unit("K"), "K": parse_unit("MPa√m"), "CGR": parse_unit("mm/s")}
    assert np.allclose(formula_dims("Temp + 273.15", units), (0, 0, 0, 1, 0))
    assert is_dimensionless("Q / (R * T_K)", units)[0]
    assert not is_dimensionless("K / CGR", units)[0]
    try:
        formula_dims("exp(K)", units)
        raise AssertionError("exp of a dimensional quantity accepted")
    except DimensionError:
        pass
    assert si_formula("CGR * 2", units) == "CGR * 0.001 * 2"
    assert si_formula("Temp / 2", units) == "(Temp + 273.15) / 2"


def test_collapse_score():
    from scc_core.collapse import collapse
    rng = np.random.default_rng(0)
    x = pd.Series(10 ** rng.uniform(0, 2, 200))
    lab = pd.Series(np.repeat(list("abcd"), 50))
    y = pd.Series(1e-7 * x ** 2.5 * 10 ** rng.normal(0, 0.2, 200))
    c = collapse(x, y, lab)
    assert abs(c.coefs[0] - 2.5) < 0.1 and abs(c.scatter - 0.2) < 0.03 and c.n_labs == 4
    # constants never change the score
    c2 = collapse(x * 123.0, y * 1e-3, lab)
    assert np.isclose(c.scatter, c2.scatter) and np.isclose(c.heldout, c2.heldout)


if __name__ == "__main__":
    for name, fn in list(globals().items()):
        if name.startswith("test_"):
            fn()
            print("ok", name)
